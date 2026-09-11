"""AI Writers' Room -- web edition.

Reuses core/ from the desktop app UNCHANGED (fdx_utils, ai_client, ai_json, xlsx_utils).
Only the GUI layer differs: CustomTkinter -> Streamlit, and the API key lives in this
app's secrets instead of each user's .env file.

Six tools are included, all sharing one pattern (load .fdx -> one whole-script AI call ->
parse JSON -> render + xlsx report). Each tool's prompt, JSON schema and report layout are
copied unchanged from the matching desktop tab in tabs/*.py, so results should match the
desktop app exactly:
  - Comedy Analysis
  - Dialogue Craft
  - Hero's Journey Editor
  - Plot Hole Editor
  - Readability Editor
  - Spelling & Grammar

Description Improver and Voice Editor are intentionally not included in this build.
"""
import datetime
import io
from pathlib import Path
import tempfile

import streamlit as st
from openpyxl import Workbook

from core import ai_client, ai_json, fdx_utils, xlsx_utils

st.set_page_config(page_title="AI Writers' Room", layout="wide")

DAILY_LIMIT = 20  # runs per browser session per day -- adjust to taste


# ======================================================================
# Per-tool configuration: prompt + JSON schema copied verbatim from the
# matching desktop tab, plus a report writer and an in-page renderer.
# ======================================================================

COMEDY_PROMPT = """You are a comedy script consultant analysing this screenplay's comedic elements scene by scene, with the goal of increasing how funny it is (smiles per minute) while staying completely faithful to the style of comedy this specific piece already uses.

FIRST: IDENTIFY THE COMEDIC STYLE

Before evaluating anything, establish what kind of comedy this screenplay actually uses -- based only on evidence in the script itself. Consider things like: wit and wordplay, irony, deadpan, character-based comedy, situational comedy, farce, satire, absurdism, black comedy, physical comedy/slapstick, running gags, comedy of manners, awkwardness/cringe comedy, and so on.

Also note what this piece clearly does NOT do -- if it never uses slapstick, never uses gross-out humour, never breaks the fourth wall, say so explicitly. Every suggestion you make later must stay consistent with the established style. Do not suggest slapstick for a script that is dry and verbal; do not suggest verbal wit for a script built on physical comedy; do not suggest a joke that would only work in a different register or genre than the one this piece actually uses.

SECOND: CATALOGUE THE COMEDY, SCENE BY SCENE

Go through the screenplay scene by scene (using the [SCENE N] markers). For each scene:

* Identify what comedic elements, if any, are present -- specific jokes, comic character behaviour, situational irony, running gags, callbacks, and so on.
* Judge whether the comedy in that scene is actually working: is it landing, partially working, or not working? If it isn't working, explain concretely why -- the joke telegraphs itself, the timing is undercut by the surrounding dialogue, the target of the joke is unclear, it repeats a joke already used elsewhere, and so on.
* Judge whether this scene is one that SHOULD carry comedy, or whether it is a deliberately dramatic/serious beat where comedy would undermine the intended emotional effect. Not every scene in a comedy needs to be funny -- forcing jokes into a scene that needs to land seriously would make the script worse, not better.
* If the comedy is weak, absent, or not working AND the scene is one that could support comedy without undermining its dramatic purpose, suggest a specific, concrete way to add or strengthen the comedy -- tied to the actual characters, situation and established comedic style of this piece. If the scene should stay dramatically serious, say so and do not force a comedy suggestion.

THIRD: OVERALL ASSESSMENT AND RECOMMENDATIONS

Summarise how well the comedy is working across the whole piece: where is it strongest, where is it weakest, and is the comedic voice consistent throughout?

Then give a short list of the most impactful, most concrete and most style-consistent changes that would raise the overall smiles-per-minute of the script without changing its comedic identity or damaging its dramatic moments.

Do not manufacture problems or force distinctiveness. A quiet, non-comedic scene that is doing its dramatic job properly is not a problem to be solved."""

COMEDY_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "comedy_style": "<the established type/voice of comedy in this piece, based on evidence>",
  "style_notes": "<what this piece does NOT do comedically -- techniques or registers that would be inconsistent with its style>",
  "scenes": [
    {"scene_number": <integer scene number, matching the [SCENE N] markers>, "comedic_elements": "<what's present, or 'None' if no comedy is attempted>", "effectiveness": "<Working well, Partially working, Not working, or N/A>", "comedy_role": "<Comedic scene, Dramatic scene -- comedy would undermine it, or Could support light comedy>", "improvement": "<concrete, style-consistent suggestion, or empty string if no change is needed>"}
  ],
  "overall_assessment": "<summary of where the comedy is strongest/weakest overall, and whether the voice is consistent>",
  "smiles_per_minute_recommendations": ["<specific, style-consistent suggestion>"]
}
The "scenes" array must contain one entry per scene in the script, in order, using the scene numbers from the [SCENE N] markers."""


def write_comedy_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    xlsx_utils.write_labelled_block(ws_overview, [
        ("Comedy style", data.get("comedy_style", "")),
        ("Style notes (what this piece avoids)", data.get("style_notes", "")),
        ("Overall assessment", data.get("overall_assessment", "")),
    ])
    xlsx_utils.set_column_widths(ws_overview, {"A": 30, "B": 100})

    ws_scenes = wb.create_sheet(title="Scenes")
    ws_scenes.append(["Scene", "Heading", "Comedic Elements", "Effectiveness", "Comedy Role", "Improvement"])
    by_scene = {}
    for entry in data.get("scenes", []):
        try:
            by_scene[int(entry.get("scene_number"))] = entry
        except (TypeError, ValueError):
            continue
    for scene_num in sorted(scene_headings):
        entry = by_scene.get(scene_num, {})
        ws_scenes.append([
            scene_num, scene_headings.get(scene_num, ""),
            entry.get("comedic_elements", ""), entry.get("effectiveness", ""),
            entry.get("comedy_role", ""), entry.get("improvement", ""),
        ])
    xlsx_utils.set_column_widths(ws_scenes, {"A": 8, "B": 30, "C": 45, "D": 18, "E": 30, "F": 55})
    xlsx_utils.enable_autofilter(ws_scenes)

    ws_recs = wb.create_sheet(title="Recommendations")
    xlsx_utils.write_list(ws_recs, "Smiles-Per-Minute Recommendations", data.get("smiles_per_minute_recommendations", []))
    xlsx_utils.set_column_widths(ws_recs, {"A": 100})

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_comedy(data: dict, scene_headings: dict) -> None:
    st.subheader("Comedy style")
    st.write(data.get("comedy_style", ""))
    st.caption(data.get("style_notes", ""))
    st.subheader("Overall assessment")
    st.write(data.get("overall_assessment", ""))
    st.subheader("Recommendations")
    for rec in data.get("smiles_per_minute_recommendations", []):
        st.markdown(f"- {rec}")
    with st.expander(f"Scene-by-scene ({len(data.get('scenes', []))} scenes)"):
        for entry in data.get("scenes", []):
            st.markdown(
                f"**Scene {entry.get('scene_number', '')} \u2014 {entry.get('effectiveness', '')}** "
                f"({entry.get('comedy_role', '')})\n\n{entry.get('comedic_elements', '')}"
            )
            if entry.get("improvement"):
                st.caption(f"Suggestion: {entry['improvement']}")
            st.divider()


# ---------------- Dialogue Craft ----------------

DIALOGUE_PROMPT = """You are a dialogue craft editor reviewing this screenplay's dialogue for problems that make lines feel written rather than spoken, or that waste the opportunity for subtext -- independent of whether any individual line matches a specific character's voice (that is a separate concern from this pass).

FIRST: ESTABLISH THE PIECE'S OWN DIALOGUE REGISTER

Before flagging anything, note how naturalistic or heightened this script's dialogue generally is -- some scripts are deliberately formal, theatrical, or stylised (a courtroom drama, a period piece, a heightened comedy), and that is a legitimate choice, not a flaw. Judge "sounds written" and "overly complete sentences" against THIS piece's own established register, not against a generic assumption that all dialogue should sound like casual conversation.

SECOND: LOOK FOR THESE PROBLEMS, LINE BY LINE

* ON-THE-NOSE -- dialogue that states exactly what the character is thinking or feeling, with no room left for the audience to read between the lines. Do not flag directness that is clearly this character's established personality -- a blunt character being blunt is not a flaw.
* UNNECESSARY GREETING OR FILLER -- greetings, acknowledgements ("Yeah." "Okay." "Right.") or repetitions that serve no dramatic, characterising or rhythmic purpose and could simply be cut.
* EXPOSITION DISGUISED AS DIALOGUE -- characters explaining plot or backstory to each other in a way that exists for the audience's benefit rather than because these characters would naturally say it to each other.
* ALREADY-KNOWN INFORMATION -- a character tells another character something both of them clearly already know, purely to inform the audience.
* OVERLY COMPLETE OR FORMAL SENTENCES -- dialogue that is grammatically tidier or more formal than how a person would actually speak, relative to this piece's own register.
* SOUNDS WRITTEN, NOT SPOKEN -- dialogue that reads like prose rather than something a person would say aloud, relative to this piece's own register.
* SUBTEXT OPPORTUNITY -- a line that works but states its meaning directly when the scene's context would let the same beat land more effectively if the character said something else while meaning this.
* COULD BE SHORTENED -- a line that could lose words without losing meaning, dramatic function, or characterisation.
* ANSWERS TOO DIRECTLY -- a character answers a question completely and directly when real conversation, and dramatic tension, would more often produce evasion, a non-answer, a question back, or an indirect response.

Only flag a line when the problem is genuine and fixing it would actually improve the scene. A blunt, plain, or simple line that is doing its job is not a problem -- do not flag dialogue merely because it could be phrased differently.

THIRD: LOOK ACROSS CHARACTERS

Separately from the line-by-line pass, consider whether different characters are distinguishable from each other by how they speak, or whether several characters share very similar speech patterns, vocabulary, or rhythm such that dialogue could be reassigned between them without much changing. Report this as a single overall observation, not a per-line finding.

FOURTH: OVERALL ASSESSMENT

Summarise how the dialogue is working overall: where it's sharpest, where it's flattest, and the single most valuable type of change the writer could make across a revision pass.

For every line-level issue, quote the exact original line (with speaker), explain briefly why it's a problem in this specific moment, and suggest a tighter or more natural alternative that preserves the line's meaning, plot information, and character relationships -- the same beat, better delivered, not a different beat.

Mark each finding's severity: "Significant" if fixing it would meaningfully improve the scene, or "Minor / optional" if it's a small polish rather than something that needs fixing."""

DIALOGUE_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "dialogue_register": "<how naturalistic or heightened this piece's dialogue generally is, based on evidence>",
  "findings": [
    {"scene": <integer, from the [SCENE N] marker>, "speaker": "<character name>", "issue_type": "On-the-nose, Unnecessary greeting or filler, Exposition disguised as dialogue, Already-known information, Overly complete or formal sentence, Sounds written not spoken, Subtext opportunity, Could be shortened, or Answers too directly", "original_line": "<exact quoted line>", "why_it_matters": "<brief explanation specific to this line>", "suggested_rewrite": "<tighter or more natural alternative preserving meaning and plot information>", "severity": "Significant or Minor / optional"}
  ],
  "speech_pattern_differentiation": "<one overall observation on whether characters are distinguishable from each other by how they speak, or whether some blur together>",
  "overall_assessment": "<summary of where the dialogue is sharpest/flattest, and the single most valuable revision-pass change>"
}
List findings in the order they occur in the script. Do not manufacture findings to pad the list -- a script with sharp, natural dialogue throughout should produce a short list."""


def write_dialogue_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    xlsx_utils.write_labelled_block(ws_overview, [
        ("Dialogue register", data.get("dialogue_register", "")),
        ("Speech pattern differentiation", data.get("speech_pattern_differentiation", "")),
        ("Overall assessment", data.get("overall_assessment", "")),
    ])
    xlsx_utils.set_column_widths(ws_overview, {"A": 28, "B": 100})

    ws_findings = wb.create_sheet(title="Findings")
    ws_findings.append(["Scene", "Speaker", "Issue Type", "Original Line", "Why It Matters", "Suggested Rewrite", "Severity"])
    for entry in data.get("findings", []):
        ws_findings.append([
            entry.get("scene", ""), entry.get("speaker", ""), entry.get("issue_type", ""),
            entry.get("original_line", ""), entry.get("why_it_matters", ""),
            entry.get("suggested_rewrite", ""), entry.get("severity", ""),
        ])
    xlsx_utils.set_column_widths(ws_findings, {"A": 8, "B": 18, "C": 26, "D": 40, "E": 40, "F": 40, "G": 16})
    xlsx_utils.enable_autofilter(ws_findings)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_dialogue(data: dict, scene_headings: dict) -> None:
    st.subheader("Dialogue register")
    st.write(data.get("dialogue_register", ""))
    st.subheader("Overall assessment")
    st.write(data.get("overall_assessment", ""))
    st.caption(data.get("speech_pattern_differentiation", ""))
    findings = data.get("findings", [])
    st.subheader(f"Findings ({len(findings)})")
    for entry in findings:
        st.markdown(
            f"**Scene {entry.get('scene', '')} \u2014 {entry.get('speaker', '')} "
            f"\u2014 {entry.get('issue_type', '')} ({entry.get('severity', '')})**\n\n"
            f"Original: *{entry.get('original_line', '')}*\n\n"
            f"Suggested: *{entry.get('suggested_rewrite', '')}*\n\n"
            f"{entry.get('why_it_matters', '')}"
        )
        st.divider()


# ---------------- Hero's Journey Editor ----------------

HERO_PROMPT = """You are a story structure consultant analysing this screenplay through the lens of the Hero's Journey (the Vogler/Campbell monomyth), with a focus on giving the writer specific, actionable notes -- not just describing what's there.

FIRST: IDENTIFY THE HERO

Identify the protagonist(s) of this story. Most screenplays have a single clear hero, but some are ensemble pieces with more than one protagonist, each potentially on their own journey. State who the hero(es) is/are and why.

SECOND: IDENTIFY THE FORMAT

Determine whether this screenplay is a self-contained story (a feature or a standalone episode) or reads as a pilot/episode of a series -- look for cues like an unresolved ending, a cliffhanger, a larger world clearly being set up, or a journey that obviously continues beyond this episode. State your judgment and the evidence for it.

If it is a pilot or episode:
* Analyse the hero's journey as it plays out within THIS episode, which may be a partial or deliberately compressed arc, and may deliberately end before resolution.
* Also describe the LARGER hero's journey this episode appears to be setting up across a series, and whether this episode does enough to establish that larger arc.

If it is a self-contained story, analyse the complete hero's journey within it.

THIRD: ANALYSE EACH STAGE

Assess the screenplay against these stages of the hero's journey:

1. Ordinary World -- the hero's status quo before the story disrupts it.
2. Call to Adventure -- the inciting incident that disrupts that world.
3. Refusal of the Call -- hesitation, fear, or reluctance to engage.
4. Meeting the Mentor -- guidance, tools or wisdom gained before committing.
5. Crossing the First Threshold -- the point of no return into the story's main conflict.
6. Tests, Allies, Enemies -- the hero learns the rules of this new world.
7. Approach to the Inmost Cave -- preparation before the central ordeal.
8. Ordeal -- the central crisis, the hero's greatest fear or biggest test so far.
9. Reward (Seizing the Sword) -- what the hero gains from surviving the ordeal.
10. The Road Back -- recommitting to finish the story, often with new stakes or pursuit.
11. Resurrection -- a final, climactic test that proves the hero's change.
12. Return with the Elixir -- the hero returns transformed, with something of value for their world.

For each stage:
* State whether it is present, weakly present, or absent.
* Cite where in the screenplay it occurs (or would need to occur).
* Briefly assess how effectively it functions dramatically.
* If it is absent, weak, or could be stronger, give a SPECIFIC, ACTIONABLE suggestion for how to add or strengthen it -- not a general note like "make it more compelling," but a concrete idea tied to this story's actual characters, world and plot.

Do not force every stage into the story artificially. Some stories compress, merge, or skip stages deliberately and effectively -- if a stage is genuinely and intentionally absent or merged with another for good dramatic reason, say so rather than inventing a problem. Your job is not to prove the template fits perfectly; it is to identify where the story's structure is genuinely working, and where a concrete change would make it work better.

FINALLY: OVERALL ACTIONABLE RECOMMENDATIONS

Beyond the stage-by-stage notes, give a short list of the most important, most concrete changes that would strengthen this screenplay's hero's journey overall. Prioritise specific, implementable ideas over general observations."""

HERO_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "protagonist": "<name(s) of the hero/heroes, and a brief note if this is an ensemble piece with more than one journey>",
  "format_assessment": "<'Self-contained story' or 'Series pilot/episode', with brief evidence for the judgment>",
  "stages": [
    {"stage": "<one of the 12 stage names listed above, in that order>", "presence": "Present, Weak, or Absent", "where": "<scene reference, using the [SCENE N] marker, or 'N/A' if absent>", "assessment": "<how well it functions dramatically>", "suggestion": "<specific, actionable suggestion -- empty string if the stage already works and needs no note>"}
  ],
  "series_level_arc": "<if a pilot/episode: the larger hero's journey this appears to set up across a series, and whether this episode lays enough groundwork for it. If self-contained: 'Not applicable -- self-contained story.'>",
  "overall_recommendations": ["<specific, implementable suggestion>"]
}
The "stages" array must contain all 12 stages in the order listed above, even where presence is "Absent"."""

STAGE_ORDER = [
    "Ordinary World", "Call to Adventure", "Refusal of the Call", "Meeting the Mentor",
    "Crossing the First Threshold", "Tests, Allies, Enemies", "Approach to the Inmost Cave",
    "Ordeal", "Reward (Seizing the Sword)", "The Road Back", "Resurrection", "Return with the Elixir",
]


def write_hero_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    xlsx_utils.write_labelled_block(ws_overview, [
        ("Protagonist(s)", data.get("protagonist", "")),
        ("Format assessment", data.get("format_assessment", "")),
        ("Series-level arc", data.get("series_level_arc", "")),
    ])
    xlsx_utils.set_column_widths(ws_overview, {"A": 22, "B": 100})

    ws_stages = wb.create_sheet(title="Stages")
    ws_stages.append(["Stage", "Presence", "Where", "Assessment", "Suggestion"])
    by_stage = {entry.get("stage", ""): entry for entry in data.get("stages", [])}
    for stage_name in STAGE_ORDER:
        entry = by_stage.get(stage_name, {})
        ws_stages.append([
            stage_name, entry.get("presence", ""), entry.get("where", ""),
            entry.get("assessment", ""), entry.get("suggestion", ""),
        ])
    xlsx_utils.set_column_widths(ws_stages, {"A": 28, "B": 12, "C": 22, "D": 50, "E": 55})
    xlsx_utils.enable_autofilter(ws_stages)

    ws_recs = wb.create_sheet(title="Recommendations")
    xlsx_utils.write_list(ws_recs, "Overall Recommendations", data.get("overall_recommendations", []))
    xlsx_utils.set_column_widths(ws_recs, {"A": 100})

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_hero(data: dict, scene_headings: dict) -> None:
    st.subheader("Protagonist(s)")
    st.write(data.get("protagonist", ""))
    st.caption(data.get("format_assessment", ""))
    if data.get("series_level_arc"):
        st.info(data["series_level_arc"])
    st.subheader("Recommendations")
    for rec in data.get("overall_recommendations", []):
        st.markdown(f"- {rec}")
    by_stage = {entry.get("stage", ""): entry for entry in data.get("stages", [])}
    with st.expander("Stage-by-stage breakdown"):
        for stage_name in STAGE_ORDER:
            entry = by_stage.get(stage_name, {})
            st.markdown(f"**{stage_name} \u2014 {entry.get('presence', '')}** ({entry.get('where', '')})")
            st.write(entry.get("assessment", ""))
            if entry.get("suggestion"):
                st.caption(f"Suggestion: {entry['suggestion']}")
            st.divider()


# ---------------- Plot Hole Editor ----------------

PLOTHOLE_PROMPT = """You are a forensic screenplay analyst. Your task is to read the screenplay as a hostile but fair-minded continuity editor, story editor and logic investigator whose primary objective is to identify anything that could make an intelligent viewer stop and think:

"Hang on -- that doesn't make sense."

Do not rewrite the screenplay and do not suggest improvements merely for the sake of improvement. Your first responsibility is to identify genuine problems, potential problems, unexplained elements and moments that may feel jarring, contrived or implausible.

Analyse the screenplay in its entirety before reaching conclusions. Do not assess scenes in isolation. Track what happens, what characters know, what characters believe, what information the audience has been given, and how events cause subsequent events.

PLOT HOLES

Look aggressively for:

* Events that cannot logically happen given what has previously been established.
* Characters doing things they would have no logical reason to do.
* Characters failing to do things they obviously would do.
* Information appearing from nowhere.
* Information being forgotten when a character should reasonably remember it.
* Problems that could easily be solved by an obvious action that nobody takes.
* Coincidences that are doing too much work.
* Characters arriving at conclusions without sufficient evidence.
* Investigations that progress because the plot requires them to rather than because of believable discoveries.
* Characters knowing things they have no plausible way of knowing.
* Characters not knowing things they have already been shown or told.
* Objects, documents, evidence, messages or information being used inconsistently.
* Events that depend on something that has not actually happened.
* Consequences that should logically follow from an event but never occur.
* Events whose consequences occur without the necessary preceding cause.

For every suspected plot hole, distinguish between: Definite plot hole, Plausible issue, and Not actually a problem.

Do not manufacture problems merely because something is not explicitly explained."""

PLOTHOLE_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "overall_assessment": "<brief summary of how tight or loose the plot logic is overall>",
  "findings": [
    {"where": "<scene reference, using the [SCENE N] marker>", "issue": "<what the potential problem is>", "verdict": "Definite plot hole, Plausible issue, or Not actually a problem", "explanation": "<why, referencing what was established elsewhere in the script>"}
  ]
}
List findings in the order they'd be noticed while reading. Include genuine "Not actually a problem" entries where something might look like a hole at first glance but is actually accounted for -- that's useful too. Do not manufacture findings just to fill the list."""


def write_plothole_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.append(["Overall assessment", data.get("overall_assessment", "")])
    xlsx_utils.set_column_widths(ws_overview, {"A": 22, "B": 100})

    ws_findings = wb.create_sheet(title="Findings")
    ws_findings.append(["Where", "Issue", "Verdict", "Explanation"])
    for entry in data.get("findings", []):
        ws_findings.append([entry.get("where", ""), entry.get("issue", ""), entry.get("verdict", ""), entry.get("explanation", "")])
    xlsx_utils.set_column_widths(ws_findings, {"A": 25, "B": 45, "C": 22, "D": 60})
    xlsx_utils.enable_autofilter(ws_findings)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_plothole(data: dict, scene_headings: dict) -> None:
    st.subheader("Overall assessment")
    st.write(data.get("overall_assessment", ""))
    findings = data.get("findings", [])
    st.subheader(f"Findings ({len(findings)})")
    for entry in findings:
        st.markdown(f"**{entry.get('where', '')} \u2014 {entry.get('verdict', '')}**\n\n{entry.get('issue', '')}")
        st.caption(entry.get("explanation", ""))
        st.divider()


# ---------------- Readability Editor ----------------

READABILITY_PROMPT = """You are a professional screenplay reader encountering this screenplay for the first time.

Read the screenplay as a genuine reader, not as an editor, proofreader or script doctor.

Your task is to report your subjective reading experience: where the screenplay engaged you, lost you, confused you, surprised you or made you want to keep reading.

Do not rewrite the screenplay.

Do not look for plot holes unless they directly affected your reading experience.

Your primary job is to describe honestly what it was like to read the script. Where you identify a genuine readability problem, you will also suggest a concrete way to address it -- see SUGGESTING FIXES near the end -- but do not let the desire to have a fix to offer change how honestly you report the problem itself, and do not suggest fixes for things that are already working.

## IMPORTANT: READ IT AS A FIRST-TIME READER

Approach the screenplay without hindsight.

At each point in the screenplay, assess your reaction based only on what the reader knows at that moment.

Do not criticise an earlier scene for failing to explain something that is intentionally and effectively revealed later.

However, if you were genuinely confused, bored or disengaged while reading, report that reaction even if the screenplay later explains the issue.

Distinguish between:

* temporary productive confusion that creates intrigue
* unproductive confusion that interrupts the reading experience
* information deliberately withheld
* information that appears to be missing or unclear

Your job is to simulate the experience of reading the screenplay for the first time, page by page.

## TRACK YOUR READING EXPERIENCE

As you read, pay particular attention to the following:

### ATTENTION

Where did your attention increase? Identify moments where you became particularly engaged, alert or eager to know what happened next. Explain briefly what created that engagement.

### LOSS OF ATTENTION

Where did your attention decrease? Identify scenes, passages or sequences where your interest weakened. Consider whether this was caused by repetition, slow pacing, too much information, lack of conflict, lack of clarity, predictable action, excessive description, scenes continuing after their dramatic purpose is complete, characters discussing rather than doing, or insufficient stakes or curiosity. Do not assume that a quiet scene is automatically a weak scene.

### CONFUSION

Where were you confused? Identify the exact point where confusion occurred. Explain what information you believed you needed at that moment.

Distinguish between INTRIGUING CONFUSION (you did not yet understand something, but wanted to find out) and DISTRACTING CONFUSION (you were unsure what was happening, who was involved, where you were, why something happened or how information connected).

### SKIMMING

Where did you feel tempted to skim? Identify passages that felt dense, repetitive, over-explained or insufficiently engaging -- long blocks of description, repetitive dialogue, exposition, scenes with little change, unnecessary detail, dialogue that repeats what the reader already knows. Do not assume that longer passages are automatically weaker.

### CURIOSITY

Where did you become curious? Identify questions the screenplay successfully created, mysteries that interested you, characters you wanted to understand, situations you wanted resolved, information you wanted to discover. Explain what made the curiosity effective.

### DISENGAGEMENT

Where did you stop caring, or come close to stopping caring? Be honest. Identify the point and explain why -- unclear stakes, insufficient emotional connection, repetition, predictability, passive characters, events without consequences, too many characters or storylines, scenes that appear disconnected from the main story. Do not manufacture disengagement. If the screenplay held your attention, say so.

### CHARACTERS

Which characters interested you? Explain when they first caught your attention, what made you curious about them, whether your interest increased or decreased, whether you ever became confused about their motivations or role. Also identify significant characters who failed to generate much interest and explain why. Do not confuse "unlikeable" with "uninteresting."

### SCENES THAT FELT LONG

Identify scenes that felt longer than their actual dramatic content justified. Ask: could I feel the scene continuing after I understood its purpose? Consider repeated beats, repeated information, delayed decisions, dialogue going in circles, unnecessary entrances or exits, scenes starting too early, scenes ending too late. Do not suggest cuts unless asked.

### PREDICTABILITY

Where did you anticipate what would happen? Distinguish GOOD ANTICIPATION (the screenplay deliberately created suspense because you feared or expected something) from PREDICTABILITY (you guessed the development because the screenplay made it too obvious or followed a familiar pattern). Explain the difference.

### SURPRISE

Where did something genuinely surprise you -- revelations, decisions, reversals, discoveries, character behaviour, plot developments? Assess whether the surprise felt EARNED (surprising but believable in retrospect), UNEXPECTED BUT PLAUSIBLE (a genuine surprise that still worked), or ARBITRARY (surprising because the screenplay had not adequately prepared the reader).

## PAGE-TURNING MOMENTUM

Throughout the screenplay, repeatedly ask: do I want to keep reading? Pay particular attention to scene openings, scene endings, major transitions, revelations, turning points, act breaks. Identify moments where your desire to continue reading INCREASED, REMAINED STRONG, WEAKENED, or DROPPED SIGNIFICANTLY. Do not judge individual scenes in isolation -- consider how each scene affects momentum into the next.

## AVOID HINDSIGHT BIAS

Do not use knowledge gained later in the screenplay when describing your reaction to an earlier scene. Distinguish "I was intrigued because I didn't understand X and wanted to know more" from "I was confused because I couldn't understand what was happening or why it mattered." This distinction is extremely important.

## SUGGESTING FIXES

For every genuine readability problem you identify -- a loss of attention, distracting confusion, a scene tempting you to skim, or disengagement -- also suggest ONE concrete way to address it, following these rules:

* Stay completely consistent with the world, characters, tone and plot already established. Do not suggest a different story, a new character, a different ending, or a different genre.
* Keep the fix small and targeted -- a specific line to cut, a specific piece of information to move earlier or later, a specific beat to trim, a specific reason to add for why a character does something, a specific repeated detail to remove. Do not suggest wholesale rewrites, restructuring an act, or generic notes like "tighten this" or "punch up the dialogue."
* Be concrete enough that the writer could apply it directly without further interpretation. "This scene runs long -- cut John's second explanation of the plan partway through, since Mary already understood it the first time" is a suggestion. "Tighten this scene" is not.
* Do not suggest a fix for something that isn't actually a problem -- productive confusion, deliberate suspense, and a quiet-but-effective scene do not need fixing, and should not get a suggestion.

## IMPORTANT RULES

* Read as a first-time reader, not an editor.
* Report genuine reactions, not theoretical screenplay advice.
* Do not manufacture criticism.
* Do not assume something is a problem simply because it is unconventional.
* Do not confuse ambiguity with confusion.
* Do not confuse a quiet scene with a boring scene.
* Do not confuse an unlikeable character with an uninteresting character.
* Do not confuse predictable suspense with a predictable plot.
* Do not rewrite dialogue or action lines.
* Only suggest a fix where you have identified a genuine problem -- never for a moment that is already working.
* Focus on what actually affected your desire to continue reading.

The purpose of this analysis is to answer one question: what is it genuinely like to read this screenplay for the first time?"""

READABILITY_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "overall_reading_experience": "<concise summary: did you want to keep reading, where was it strongest, where did it lose momentum, were you generally engaged>",
  "reading_journey": [
    {"where": "<scene/page/sequence, referencing the [SCENE N] marker>", "reaction": "<Engaged, Curious, Confused, Surprised, Attention dropped, Tempted to skim, or Disengaged>", "what_happened": "<brief description>", "why": "<why it affected the reading experience>", "suggestion": "<concrete, small, world-consistent fix if this entry is a genuine problem -- Attention dropped, Tempted to skim, Disengaged, or Confused-in-the-distracting-sense; empty string for positive reactions or productive/intentional confusion>"}
  ],
  "attention_high_points": ["<moment the screenplay was most compelling>"],
  "attention_low_points": [
    {"point": "<moment engagement weakened -- omit trivial criticisms>", "suggestion": "<concrete, small, world-consistent fix>"}
  ],
  "confusion_vs_curiosity": {
    "productive_mysteries": "<intriguing unanswered questions that worked>",
    "distracting_confusion": [
      {"issue": "<confusion that interrupted the reading experience, and why it was distracting rather than intriguing>", "suggestion": "<concrete, small, world-consistent fix>"}
    ],
    "information_overload": "<any -- or state none>"
  },
  "character_engagement": {
    "most_interesting": "<characters and why>",
    "increased_interest": "<characters whose interest grew, and when>",
    "decreased_interest": "<characters whose interest faded, and when>",
    "failed_to_impress": "<significant characters who made little impression, and why>"
  },
  "pacing_and_skimming": {
    "felt_long_or_repetitive": [
      {"issue": "<scene/passage that felt too long, slow, dense or over-explained>", "suggestion": "<concrete, small, world-consistent fix>"}
    ],
    "moved_well": "<scenes/passages that moved particularly well>"
  },
  "predictability_and_surprise": {
    "anticipated": "<developments you guessed, and whether that was good anticipation or predictability>",
    "genuine_surprises": "<earned or unexpected-but-plausible surprises>",
    "arbitrary_surprises": "<surprises that felt unprepared for, if any>"
  },
  "final_verdict": {
    "would_keep_reading": "<yes/no and why>",
    "likely_stop_point": "<where you'd stop if other scripts were competing for attention>",
    "three_strongest_elements": "<three strongest elements of the reading experience>",
    "three_biggest_threats": "<three biggest threats to reader engagement>"
  }
}
List reading_journey in chronological order, one entry per significant change in reaction -- not one per scene. attention_low_points, distracting_confusion, and felt_long_or_repetitive are empty arrays [] if there's nothing genuine to report -- do not pad them, and do not invent a suggestion for an entry that isn't really a problem."""


def write_readability_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.append(["Overall reading experience", data.get("overall_reading_experience", "")])
    verdict = data.get("final_verdict", {})
    xlsx_utils.write_labelled_block(ws_overview, [
        ("Would keep reading?", verdict.get("would_keep_reading", "")),
        ("Likely stop point", verdict.get("likely_stop_point", "")),
        ("Three strongest elements", verdict.get("three_strongest_elements", "")),
        ("Three biggest threats", verdict.get("three_biggest_threats", "")),
    ])
    ws_overview.column_dimensions["A"].width = 28
    ws_overview.column_dimensions["B"].width = 100

    ws_journey = wb.create_sheet(title="Reading Journey")
    ws_journey.append(["Where", "Reaction", "What Happened", "Why", "Suggestion"])
    for entry in data.get("reading_journey", []):
        ws_journey.append([
            entry.get("where", ""), entry.get("reaction", ""), entry.get("what_happened", ""),
            entry.get("why", ""), entry.get("suggestion", ""),
        ])
    for col, width in zip("ABCDE", (30, 18, 40, 45, 55)):
        ws_journey.column_dimensions[col].width = width
    xlsx_utils.enable_autofilter(ws_journey)

    ws_attention = wb.create_sheet(title="Attention & Pacing")
    ws_attention.append(["Attention High Points"])
    for item in data.get("attention_high_points", []):
        ws_attention.append([item])
    ws_attention.append([])
    ws_attention.append(["Attention Low Points", "Suggestion"])
    for item in data.get("attention_low_points", []):
        ws_attention.append([item.get("point", ""), item.get("suggestion", "")])
    ws_attention.append([])
    ws_attention.append(["Felt Long / Repetitive", "Suggestion"])
    pacing = data.get("pacing_and_skimming", {})
    for item in pacing.get("felt_long_or_repetitive", []):
        ws_attention.append([item.get("issue", ""), item.get("suggestion", "")])
    ws_attention.append([])
    ws_attention.append(["Moved well", pacing.get("moved_well", "")])
    ws_attention.column_dimensions["A"].width = 55
    ws_attention.column_dimensions["B"].width = 55

    ws_deep = wb.create_sheet(title="Confusion Chars Surprise")
    confusion = data.get("confusion_vs_curiosity", {})
    characters = data.get("character_engagement", {})
    surprise = data.get("predictability_and_surprise", {})
    xlsx_utils.write_labelled_block(ws_deep, [
        ("Productive mysteries", confusion.get("productive_mysteries", "")),
        ("Information overload", confusion.get("information_overload", "")),
    ])
    ws_deep.append([])
    ws_deep.append(["Distracting Confusion", "Suggestion"])
    for item in confusion.get("distracting_confusion", []):
        ws_deep.append([item.get("issue", ""), item.get("suggestion", "")])
    ws_deep.append([])
    xlsx_utils.write_labelled_block(ws_deep, [
        ("Most interesting characters", characters.get("most_interesting", "")),
        ("Interest increased", characters.get("increased_interest", "")),
        ("Interest decreased", characters.get("decreased_interest", "")),
        ("Failed to impress", characters.get("failed_to_impress", "")),
        ("Anticipated developments", surprise.get("anticipated", "")),
        ("Genuine surprises", surprise.get("genuine_surprises", "")),
        ("Arbitrary surprises", surprise.get("arbitrary_surprises", "")),
    ])
    ws_deep.column_dimensions["A"].width = 28
    ws_deep.column_dimensions["B"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_readability(data: dict, scene_headings: dict) -> None:
    st.subheader("Overall reading experience")
    st.write(data.get("overall_reading_experience", ""))
    verdict = data.get("final_verdict", {})
    st.markdown(f"**Would keep reading?** {verdict.get('would_keep_reading', '')}")
    st.caption(f"Likely stop point: {verdict.get('likely_stop_point', '')}")
    st.markdown(f"**Strongest elements:** {verdict.get('three_strongest_elements', '')}")
    st.markdown(f"**Biggest threats:** {verdict.get('three_biggest_threats', '')}")
    with st.expander(f"Reading journey ({len(data.get('reading_journey', []))} entries)"):
        for entry in data.get("reading_journey", []):
            st.markdown(f"**{entry.get('where', '')} \u2014 {entry.get('reaction', '')}**\n\n{entry.get('what_happened', '')}")
            st.caption(entry.get("why", ""))
            if entry.get("suggestion"):
                st.caption(f"Suggestion: {entry['suggestion']}")
            st.divider()


# ---------------- Spelling & Grammar ----------------

SPELLING_PROMPT = """You are a meticulous copy editor proofreading this screenplay for misspellings, missing words, and repeated words. Your job is mechanical accuracy, not creative or story feedback -- do not comment on dialogue quality, characterisation, or story logic.

The screenplay is tagged by paragraph type: [ACTION] for narrative description, [DIALOGUE:CHARACTER] for spoken lines, [CHARACTER] for character cues, [PARENTHETICAL] for parentheticals, [TRANSITION] for transitions. Apply different standards to each:

ACTION lines are narrative prose and should be grammatically complete and correctly spelled. Flag genuine errors here fairly readily.

DIALOGUE is spoken by characters and is often deliberately fragmented, incomplete, interrupted, or written in a particular character's natural speech pattern -- dropped words, trailing off, informal grammar, dialect spelling like "gonna" or "ain't". This is a deliberate writing choice, NOT an error -- do not flag it. Only flag something in dialogue when it reads as an accidental slip of the pen rather than a natural speech pattern -- most reliably, a word obviously repeated back-to-back (e.g. "to to", "the the", "was was"), not a stylistic incompleteness.

Do not check [SCENE N] markers, scene heading text, or [CHARACTER] cue lines themselves -- focus on ACTION, DIALOGUE, and PARENTHETICAL content.

WHAT TO LOOK FOR

1. MISSPELLINGS -- genuine spelling errors. Do not flag:
   * Character names, place names, or invented words used consistently throughout the script -- these are intentional, not typos.
   * Screenplay format terms (INT., EXT., O.S., V.O., CONT'D, etc).
   * Deliberate dialect or phonetic spelling in dialogue (e.g. "gonna", "ain't", "y'know").

2. MISSING WORDS -- a short word (usually an article, preposition, or conjunction) appears to have been dropped, making the sentence read as broken rather than intentionally terse. For example, "She went the hospital" is clearly missing "to". Flag these readily in ACTION lines. In DIALOGUE, only flag if it reads as an accidental gap rather than a natural way a character might actually speak.

3. REPEATED WORDS -- the same short word appears twice in a row by accident, e.g. "He walks over to to the door" or "the the door". These are almost always genuine slips regardless of whether they appear in action or dialogue, and should be flagged confidently. Do not confuse this with a deliberate stutter written with punctuation (e.g. "I... I didn't mean it"), which is a style choice, not an error.

For each issue found, quote the exact short phrase containing it (not the whole sentence) so it can be located and corrected precisely, and give the corrected version of that exact phrase.

Distinguish your confidence:
* "Definite error" -- unambiguous, no reasonable reading where this is intentional.
* "Likely error" -- very probably a mistake, but a small chance it's deliberate.
* "Possibly intentional -- flagged for awareness" -- could go either way; flagging only so the writer can make the call themselves.

Do not manufacture findings to pad the list. A screenplay with clean prose and naturalistic dialogue should produce a short list, not a long one."""

SPELLING_JSON = """

Return ONLY valid JSON, no markdown formatting, no commentary before or after, matching exactly this schema:
{
  "findings": [
    {"scene": <integer, from the [SCENE N] marker>, "context_type": "Action, Dialogue, Parenthetical, or Other", "speaker": "<character name if Dialogue, otherwise empty string>", "issue_type": "Misspelling, Missing word, or Repeated word", "original_text": "<short exact phrase from the script containing the issue>", "suggested_fix": "<corrected version of that exact phrase>", "verdict": "Definite error, Likely error, or Possibly intentional -- flagged for awareness"}
  ],
  "overall_note": "<brief note on any recurring pattern noticed, e.g. a word the writer often drops -- or an empty string if nothing notable>"
}
List findings in the order they occur in the script. It is fine for this list to be short or even empty if the script is clean."""


def write_spelling_report(data: dict, scene_headings: dict) -> io.BytesIO:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.append(["Overall note", data.get("overall_note", "")])
    xlsx_utils.set_column_widths(ws_overview, {"A": 18, "B": 100})

    ws_findings = wb.create_sheet(title="Findings")
    ws_findings.append(["Scene", "Context", "Speaker", "Issue Type", "Original", "Suggested Fix", "Verdict"])
    for entry in data.get("findings", []):
        ws_findings.append([
            entry.get("scene", ""), entry.get("context_type", ""), entry.get("speaker", ""),
            entry.get("issue_type", ""), entry.get("original_text", ""),
            entry.get("suggested_fix", ""), entry.get("verdict", ""),
        ])
    xlsx_utils.set_column_widths(ws_findings, {"A": 8, "B": 12, "C": 18, "D": 15, "E": 40, "F": 40, "G": 32})
    xlsx_utils.enable_autofilter(ws_findings)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def render_spelling(data: dict, scene_headings: dict) -> None:
    if data.get("overall_note"):
        st.info(data["overall_note"])
    findings = data.get("findings", [])
    st.subheader(f"Findings ({len(findings)})")
    if not findings:
        st.write("No issues found.")
    for entry in findings:
        st.markdown(
            f"**Scene {entry.get('scene', '')} \u2014 {entry.get('issue_type', '')} "
            f"({entry.get('verdict', '')})**\n\n"
            f"`{entry.get('original_text', '')}` \u2192 `{entry.get('suggested_fix', '')}`"
        )
        st.divider()


# ======================================================================
# Tool registry
# ======================================================================

TOOLS = {
    "Comedy Analysis": {
        "prompt": COMEDY_PROMPT, "json_instructions": COMEDY_JSON,
        "extract_fn": fdx_utils.extract_full_script,
        "write_report": write_comedy_report, "render": render_comedy,
        "file_suffix": "comedy",
    },
    "Dialogue Craft": {
        "prompt": DIALOGUE_PROMPT, "json_instructions": DIALOGUE_JSON,
        "extract_fn": fdx_utils.extract_tagged_script,
        "write_report": write_dialogue_report, "render": render_dialogue,
        "file_suffix": "dialogue_craft",
    },
    "Hero's Journey Editor": {
        "prompt": HERO_PROMPT, "json_instructions": HERO_JSON,
        "extract_fn": fdx_utils.extract_full_script,
        "write_report": write_hero_report, "render": render_hero,
        "file_suffix": "heros_journey",
    },
    "Plot Hole Editor": {
        "prompt": PLOTHOLE_PROMPT, "json_instructions": PLOTHOLE_JSON,
        "extract_fn": fdx_utils.extract_full_script,
        "write_report": write_plothole_report, "render": render_plothole,
        "file_suffix": "plotholes",
    },
    "Readability Editor": {
        "prompt": READABILITY_PROMPT, "json_instructions": READABILITY_JSON,
        "extract_fn": fdx_utils.extract_full_script,
        "write_report": write_readability_report, "render": render_readability,
        "file_suffix": "readability",
    },
    "Spelling & Grammar": {
        "prompt": SPELLING_PROMPT, "json_instructions": SPELLING_JSON,
        "extract_fn": fdx_utils.extract_tagged_script,
        "write_report": write_spelling_report, "render": render_spelling,
        "file_suffix": "spellcheck",
    },
}


# ======================================================================
# Access gate
# ======================================================================

def check_access() -> bool:
    if st.session_state.get("authenticated"):
        return True
    st.title("AI Writers' Room")
    st.caption("Friends & testers preview")
    with st.form("login"):
        name = st.text_input("Your name (so Steve can tell whose run is whose)")
        password = st.text_input("Password", type="password")
        submitted = st.form_submit_button("Enter")
    if submitted:
        if not name.strip():
            st.error("Please enter your name.")
        elif password == st.secrets.get("APP_PASSWORD"):
            st.session_state["authenticated"] = True
            st.session_state["user_name"] = name.strip()
            st.rerun()
        else:
            st.error("Incorrect password.")
    return False


if not check_access():
    st.stop()

# ---------- Usage cap (per browser session, resets daily) ----------
today = str(datetime.date.today())
if st.session_state.get("usage_date") != today:
    st.session_state["usage"] = 0
    st.session_state["usage_date"] = today


def credits_left() -> int:
    return DAILY_LIMIT - st.session_state["usage"]


def use_credit() -> None:
    st.session_state["usage"] += 1


# ---------- API key comes from app secrets, never from the user ----------
API_KEY = st.secrets["ANTHROPIC_API_KEY"]
MODEL = st.secrets.get("ANTHROPIC_MODEL", "claude-sonnet-4-6")
MAX_OUTPUT_TOKENS = int(st.secrets.get("MAX_OUTPUT_TOKENS", 16000))


# ---------- Optional run logging to a Google Sheet ----------
def log_run(fdx_filename: str, mode: str, status: str, error_msg: str = "", notes: str = "") -> None:
    if "gcp_service_account" not in st.secrets:
        st.sidebar.warning("DEBUG: gcp_service_account not found in secrets at all.")
        return
    try:
        import gspread
        from google.oauth2.service_account import Credentials

        creds = Credentials.from_service_account_info(
            st.secrets["gcp_service_account"],
            scopes=[
                "https://www.googleapis.com/auth/spreadsheets",
                "https://www.googleapis.com/auth/drive",
            ],
        )
        gc = gspread.authorize(creds)
        sheet = gc.open(st.secrets.get("LOG_SHEET_NAME", "AI Writers Room Logs")).sheet1
        sheet.append_row([
            datetime.datetime.now().isoformat(),
            st.session_state.get("user_name", "unknown"),
            fdx_filename, mode, status, error_msg, notes,
        ])
    except Exception as e:
        st.sidebar.warning(f"DEBUG: logging failed -- {type(e).__name__}: {e}")


def save_upload_to_tempfile(uploaded_file) -> str:
    suffix = Path(uploaded_file.name).suffix or ".fdx"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(uploaded_file.getbuffer())
        return tmp.name


# ======================================================================
# Sidebar + main flow
# ======================================================================

st.sidebar.write(f"Signed in as **{st.session_state['user_name']}**")
st.sidebar.write(f"Runs left today: **{credits_left()}**")
mode = st.sidebar.selectbox("Choose a tool", list(TOOLS.keys()))

st.title(f"AI Writers' Room \u2014 {mode}")

if credits_left() <= 0:
    st.error("Daily usage limit reached for this session. Please try again tomorrow.")
    st.stop()

cfg = TOOLS[mode]
prompt = st.text_area("Analysis instructions", value=cfg["prompt"], height=220)
uploaded = st.file_uploader("Upload a .fdx file", type=["fdx"], key=f"upload_{mode}")

run_key = f"result_{mode}"

if uploaded and st.button("Run analysis"):
    fdx_path = save_upload_to_tempfile(uploaded)
    try:
        _tree, root = fdx_utils.load_fdx(fdx_path)
    except Exception as e:
        st.error(f"Could not open file: {e}")
        log_run(uploaded.name, mode, "error", str(e))
        st.stop()

    script_text, scene_headings = cfg["extract_fn"](root)
    if not script_text.strip():
        st.warning("No script content found.")
        st.stop()

    with st.spinner(f"Analysing {len(scene_headings)} scenes..."):
        try:
            raw = ai_client.send_prompt(
                API_KEY, MODEL,
                system=prompt + cfg["json_instructions"],
                user_content=script_text,
                max_tokens=MAX_OUTPUT_TOKENS,
            )
            data = ai_json.parse_json_response(raw)
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            log_run(uploaded.name, mode, "error", str(e))
            st.stop()

    use_credit()
    log_run(uploaded.name, mode, "success")
    st.session_state[run_key] = {
        "data": data, "scene_headings": scene_headings, "fdx_name": uploaded.name,
    }

if st.session_state.get(run_key):
    result = st.session_state[run_key]
    cfg["render"](result["data"], result["scene_headings"])
    buf = cfg["write_report"](result["data"], result["scene_headings"])
    st.download_button(
        "Download report (.xlsx)",
        data=buf.getvalue(),
        file_name=Path(result["fdx_name"]).stem + f"_{cfg['file_suffix']}.xlsx",
    )
